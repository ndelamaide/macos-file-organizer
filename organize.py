# /// script
# requires-python = ">=3.11"
# dependencies = [
#     "ocrmac",
#     "ollama",
#     "pillow-heif",
#     "pydantic",
#     "pypdf",
#     "pypdfium2",
#     "python-docx",
# ]
# ///
"""
Local AI file organizer (Ollama). Nothing leaves your machine.

  uv run organize.py plan ~/Downloads ~/Desktop "~/My Drive"   # writes runs/plan-<timestamp>.csv, moves nothing
  uv run organize.py plan "~/My Drive" --refresh                # also re-date/rename files already sorted
  # -> review the plan: edit dest/category, set action to "skip" for rows to leave alone
  uv run organize.py apply                                      # executes the latest plan, logs runs/undo-<timestamp>.csv
  uv run organize.py apply runs/plan-20260926-165700.csv        # or a specific plan

Settings live in config.toml next to this script (start from config.example.toml).
"""
import argparse, csv, difflib, hashlib, os, re, shutil, string, sys, tomllib, unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Literal

import ollama
from pydantic import create_model

CONFIG_PATH = Path(__file__).with_name("config.toml")
RUNS_DIR = Path(__file__).with_name("runs")  # plans and undo logs, one file per run
PLAN_FIELDS = ["action", "src", "dest", "category", "group", "subfolder", "date",
               "suggested", "confidence", "note"]

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".heic", ".tiff", ".tif"}
TEXT_EXT = {".txt", ".md", ".csv", ".json"}
GOOGLE_STUBS = {".gdoc", ".gsheet", ".gslides", ".gdraw", ".gform", ".gmap", ".gsite"}
BUNDLES = (".app", ".photoslibrary", ".bundle", ".framework", ".pkg")
MIN_PDF_TEXT = 50  # below this many characters, treat the PDF as a scan and OCR it

DEFAULT_DATE_RULE = ("the date the document is about or was issued (e.g. invoice date, "
                     "letter date, statement period)")
DEFAULT_FILENAME = "{date}-{title}"
PRECISIONS = ("year", "month", "day")
TEMPLATE_FIELDS = {"date", "title", "group", "subfolder"}


# ---- config --------------------------------------------------------------

def expand(p: str) -> Path:
    return Path(os.path.expanduser(p)).resolve()


def naming_settings(v: dict, where: str) -> dict:
    """date / date_precision / filename, as set on a category or a subfolder."""
    s = {"date": v.get("date"), "date_precision": v.get("date_precision"), "filename": v.get("filename")}
    if s["date_precision"] and s["date_precision"] not in PRECISIONS:
        sys.exit(f"config.toml: {where}: date_precision must be one of {', '.join(PRECISIONS)}")
    if s["filename"]:
        used = {f for _, f, _, _ in string.Formatter().parse(s["filename"]) if f}
        if used - TEMPLATE_FIELDS:
            sys.exit(f"config.toml: {where}: unknown placeholder(s) {used - TEMPLATE_FIELDS} in filename; "
                     f"use {', '.join('{' + f + '}' for f in sorted(TEMPLATE_FIELDS))}")
    return s


def load_config() -> dict:
    if not CONFIG_PATH.exists():
        sys.exit(f"No config at {CONFIG_PATH}\nCreate it with: cp config.example.toml config.toml")
    with open(CONFIG_PATH, "rb") as fh:
        raw = tomllib.load(fh)

    cfg = {
        "model": raw.get("model", "qwen3.5:9b"),
        "max_chars": raw.get("max_chars", 3000),
        "ocr_languages": raw.get("ocr_languages", ["fr-FR", "en-US"]),
        "min_confidence": raw.get("min_confidence", 0.6),
        "inbox": raw.get("inbox", "Inbox"),
        "title_language": raw.get("title_language", ""),
        "exclude": {e for e in raw.get("exclude", []) if not e.startswith(("/", "~"))},
        "exclude_paths": {expand(e) for e in raw.get("exclude", []) if e.startswith(("/", "~"))},
        "root": expand(raw["destination"]) if "destination" in raw else None,
        "naming": naming_settings(raw, "top level"),
    }
    if "cloud" in cfg["model"]:
        sys.exit("config.toml: cloud models would send your files off-device; use a local model")
    if not raw.get("categories"):
        sys.exit("config.toml: add at least one entry under [categories]")

    cats = {}
    for name, v in raw["categories"].items():
        if isinstance(v, str):
            v = {"description": v}
        subs_raw = v.get("subfolders") or {}
        if isinstance(subs_raw, list):  # ["Contracts", "Payslips"]
            subs_raw = {s: {} for s in subs_raw}
        subs = {}
        for s, sv in subs_raw.items():  # Payslips = "when to use it"  or  Payslips = { description, date… }
            sv = {"description": sv} if isinstance(sv, str) else sv
            subs[s] = {"description": sv.get("description", "")} | naming_settings(sv, f"{name} > {s}")
        cats[name] = {
            "description": v.get("description", ""),
            "destination": v.get("destination"),
            "group_by": v.get("group_by"),  # e.g. "employer": one folder per employer
            "subfolders": subs,             # e.g. Contracts, Payslips: one folder per document type
        } | naming_settings(v, name)
    cats.setdefault(cfg["inbox"], {"description": "No category above clearly fits", "destination": None,
                                   "group_by": None, "subfolders": {}} | naming_settings({}, "inbox"))
    cfg["categories"] = cats
    return cfg


def category_dir(cfg: dict, name: str) -> Path:
    custom = cfg["categories"][name]["destination"]
    return expand(custom) if custom else cfg["root"] / name


def resolve(cfg: dict, cat: str, sub: str, key: str):
    """Subfolder setting, else category setting, else top-level setting."""
    c = cfg["categories"][cat]
    for layer in (c["subfolders"].get(sub, {}), c, cfg["naming"]):
        if layer.get(key):
            return layer[key]
    return None


# ---- scanning ------------------------------------------------------------

def warn(err: OSError):
    print(f"  ! cannot read {err.filename}: {err.strerror}", file=sys.stderr)
    if err.errno == 1:  # EPERM: macOS privacy protection
        print("    Allow your terminal app in System Settings > Privacy & Security > "
              "Files & Folders (or Full Disk Access), then restart it.", file=sys.stderr)


def iter_files(sources: list[Path], cfg: dict, sorted_dirs: set[Path]):
    """Yield new files to sort. Skips hidden items, git repos, excluded names,
    app bundles and folders that are already sorted categories."""
    seen = set()
    for src in sources:
        found = 0
        for dirpath, dirnames, filenames in os.walk(src, onerror=warn):
            d = Path(dirpath)
            if ".git" in dirnames or ".git" in filenames or d in sorted_dirs or d in cfg["exclude_paths"]:
                dirnames[:] = []  # code repo, already sorted or excluded: don't descend
                continue
            dirnames[:] = [
                n for n in dirnames
                if not n.startswith(".") and n not in cfg["exclude"] and not n.endswith(BUNDLES)
            ]
            for name in filenames:
                f = d / name
                if not name.startswith(".") and f not in seen:
                    seen.add(f)
                    found += 1
                    yield f
        print(f"Scanned {src}: {found} files", file=sys.stderr)


def is_grouped(cfg: dict, cat: str) -> bool:
    c = cfg["categories"][cat]
    return bool(c["group_by"] or c["subfolders"])


def iter_sorted(sources: list[Path], cfg: dict, sorted_dirs: set[Path], refresh: bool):
    """Yield (file, category) for files already in category folders that need another pass:
    - always: files in a grouped category that aren't in their group/subfolder yet
      (e.g. Work/Employment/payslip.pdf -> Acme/Payslips/);
    - with refresh: every sorted file, to recompute its date and name.
    Only category folders inside the scanned sources are considered. Folders deeper than
    the category's structure are yours, and left alone."""
    for cat in cfg["categories"]:
        grouped = is_grouped(cfg, cat)
        if not (grouped or refresh):
            continue
        root = category_dir(cfg, cat)
        if not root.is_dir() or not any(root.is_relative_to(s) for s in sources):
            continue
        c = cfg["categories"][cat]
        depth = bool(c["group_by"]) + bool(c["subfolders"])
        found = 0
        for dirpath, dirnames, filenames in os.walk(root, onerror=warn):
            d = Path(dirpath)
            if ".git" in dirnames or d in cfg["exclude_paths"] or (d != root and d in sorted_dirs):
                dirnames[:] = []
                continue
            level = len(d.relative_to(root).parts)
            if level >= depth:
                dirnames[:] = []  # don't look deeper than the category's own structure
            else:
                dirnames[:] = [n for n in dirnames if not n.startswith(".") and n not in cfg["exclude"]]
            if level == depth and not refresh:
                continue  # already in place
            for name in filenames:
                if not name.startswith("."):
                    found += 1
                    yield d / name, cat
        if found:
            print(f"Re-checking {root}: {found} sorted files", file=sys.stderr)


class DuplicateFinder:
    """Finds files with identical content. Compares sizes first and only
    hashes (SHA-256) when sizes match, so large trees stay fast."""

    def __init__(self):
        self.by_size: dict[int, list[Path]] = {}
        self._hashes: dict[Path, str] = {}

    def _hash(self, p: Path) -> str:
        if p not in self._hashes:
            h = hashlib.sha256()
            with open(p, "rb") as fh:
                for chunk in iter(lambda: fh.read(1 << 20), b""):
                    h.update(chunk)
            self._hashes[p] = h.hexdigest()
        return self._hashes[p]

    def add(self, p: Path):
        self.by_size.setdefault(p.stat().st_size, []).append(p)

    def check_and_add(self, p: Path) -> Path | None:
        """Return an existing file with the same content, or register p."""
        for other in self.by_size.get(p.stat().st_size, []):
            if self._hash(other) == self._hash(p):
                return other
        self.add(p)
        return None


def index_sorted(sorted_dirs: set[Path], finder: DuplicateFinder):
    """Register files already in category folders, so re-runs catch copies of them."""
    for d in sorted_dirs:
        for dirpath, dirnames, filenames in os.walk(d):
            dirnames[:] = [n for n in dirnames if not n.startswith(".")]
            for name in filenames:
                if not name.startswith("."):
                    finder.add(Path(dirpath) / name)


# ---- text extraction ------------------------------------------------------

def ocr(image, cfg: dict) -> str:
    from ocrmac import ocrmac  # Apple Vision, on-device
    res = ocrmac.OCR(image, language_preference=cfg["ocr_languages"]).recognize()
    return " ".join(text for text, _, _ in res)


def extract_text(p: Path, cfg: dict) -> str:
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(pg.extract_text() or "" for pg in PdfReader(p).pages[:3])
            if len(text.strip()) < MIN_PDF_TEXT:  # scanned PDF: OCR the first page
                import pypdfium2 as pdfium
                text = ocr(pdfium.PdfDocument(p)[0].render(scale=2).to_pil(), cfg)
            return text
        if ext in IMAGE_EXT:
            from PIL import Image
            from pillow_heif import register_heif_opener
            register_heif_opener()
            return ocr(Image.open(p), cfg)
        if ext == ".docx":
            import docx
            return "\n".join(par.text for par in docx.Document(p).paragraphs)
        if ext in TEXT_EXT:
            return p.read_text(errors="ignore")
    except Exception as e:
        print(f"  ! could not read {p.name}: {e}")
    return ""  # Google stubs, audio, etc. -> classified by filename only


# ---- model calls ------------------------------------------------------------

def ask(prompt: str, schema, cfg: dict):
    r = ollama.chat(
        model=cfg["model"],
        messages=[{"role": "user", "content": prompt}],
        format=schema.model_json_schema(),
        options={"temperature": 0},
        think=False,  # reasoning models (e.g. qwen3.5) would otherwise "think" before every answer: much slower, no gain here
    )
    return schema.model_validate_json(r.message.content)


def folder_label(p: Path, sources: list[Path]) -> str:
    """Where the file sits, from the scanned folder's name down: "Downloads/Taxes 2023"."""
    for s in sources:
        if p.is_relative_to(s):
            return str(p.parent.relative_to(s.parent))
    return p.parent.name


def file_block(p: Path, text: str, folder: str) -> str:
    return f"Filename: {p.name}\nFolder: {folder}\nContent:\n{text or '(no readable content)'}"


def make_category_model(cfg: dict, folder_check: bool = False):
    fields = {"category": (Literal[tuple(cfg["categories"])], ...), "confidence": (float, ...)}
    if folder_check:
        fields["folder_fits"] = (bool, ...)
    return create_model("Category", **fields)


FOLDER_CHECK = (
    "Also say whether the file's current folder already fits it (folder_fits). True only if that folder "
    "is a specific, meaningful home for this file: named after its project, school, topic, organization "
    "or event, and the file clearly belongs there. False for generic or catch-all folders (Downloads, "
    "Desktop, Documents, Scans, Misc, Old, Untitled, New folder) or if the file looks misplaced.\n"
)


def classify(p: Path, text: str, folder: str, cfg: dict, Category):
    """Step 1: which category, and (with a folder_fits field) whether the file can stay where it is."""
    menu = "\n".join(f"- {name}: {c['description']}" for name, c in cfg["categories"].items())
    prompt = (
        "You sort personal files (documents may be in French or English). "
        f"Pick the single best category:\n{menu}\n\n"
        "The filename and folder are hints too: the file may already be named or filed sensibly.\n"
        "Give your confidence from 0 to 1.\n"
        + (FOLDER_CHECK if "folder_fits" in Category.model_fields else "")
        + "\n" + file_block(p, text, folder)
    )
    return ask(prompt, Category, cfg)


def describe(p: Path, text: str, folder: str, cat: str, cfg: dict, known: list[str]) -> dict:
    """Step 2, now that the category is known: group, subfolder, the date that matters
    for this kind of document, and a short title."""
    c = cfg["categories"][cat]
    fields, asks = {}, []

    if c["group_by"]:
        g = c["group_by"]
        fields["group"] = (str, ...)
        asks.append(
            f"- group: which {g} is this document about? Give its usual short name, without "
            f"legal suffixes like SA, SAS, SARL, AG, GmbH, Ltd. "
            + (f"Known so far: {', '.join(known)}. If it is one of these, return exactly that name. "
               if known else "")
            + 'If you cannot tell, return "".'
        )

    if c["subfolders"]:
        fields["subfolder"] = (Literal[tuple(c["subfolders"])], ...)
        options = "\n".join(f"  - {s}" + (f": {v['description']}" if v["description"] else "")
                            for s, v in c["subfolders"].items())
        asks.append(f"- subfolder: which type of document is it?\n{options}")

    fields["date"] = (str, ...)
    base_rule = c["date"] or cfg["naming"]["date"] or DEFAULT_DATE_RULE
    sub_rules = {s: v["date"] for s, v in c["subfolders"].items() if v["date"]}
    if sub_rules:
        rule = ("depends on the subfolder:\n"
                + "\n".join(f"    - {s}: {r}" for s, r in sub_rules.items())
                + f"\n    - otherwise: {base_rule}")
    else:
        rule = base_rule
    asks.append(
        f"- date: {rule}.\n"
        "  Format YYYY-MM-DD, or YYYY-MM or YYYY if only that much is known. Read it from the content. "
        "Ignore print, download, scan and file dates, and dates in the filename unless the content "
        'confirms them. If the document does not show this date, return "".'
    )

    fields["title"] = (str, ...)
    lang = f"in {cfg['title_language']}" if cfg["title_language"] else "in the document's language"
    avoid = f", the {c['group_by']}" if c["group_by"] else ""
    asks.append(
        f"- title: 2 to 6 words saying what the document is, {lang}. If the current filename already says "
        "that clearly (e.g. \"Rapport de stage Nestlé\", not \"scan_0034\" or \"document (3)\"), reuse its "
        f"wording. No date, no file extension, don't repeat the category{avoid}. "
        "E.g. \"avis d'imposition\", \"bank statement\", \"train ticket geneva paris\"."
    )

    prompt = (
        f'This file was filed under "{cat}" ({c["description"]}). Answer:\n'
        + "\n".join(asks) + "\n\n" + file_block(p, text, folder)
    )
    r = ask(prompt, create_model("Details", **fields), cfg)

    group = ""
    if c["group_by"]:
        group = clean_group(getattr(r, "group", "") or "")
        group = match_group(group, known) if group else f"Unknown {c['group_by']}"
    sub = getattr(r, "subfolder", "") or ""
    precision = resolve(cfg, cat, sub, "date_precision") or "day"
    return {"group": group, "subfolder": sub, "date": normalize_date(r.date, precision), "title": r.title}


# ---- dates and names --------------------------------------------------------

def normalize_date(s: str, precision: str) -> str:
    """Parse what the model returned into YYYY[-MM[-DD]], truncated to `precision`.
    Anything implausible becomes "" (no date is better than a wrong one)."""
    s = (s or "").strip()
    y = m = d = None
    if mt := re.fullmatch(r"(\d{4})(?:[-/.](\d{1,2})(?:[-/.](\d{1,2}))?)?", s):
        y, m, d = mt.groups()
    elif mt := re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", s):  # 31.03.2024 (European order)
        d, m, y = mt.groups()
    elif mt := re.fullmatch(r"(\d{1,2})[-/.](\d{4})", s):  # 03/2024
        m, y = mt.groups()
    if y is None:
        return ""
    y, m, d = int(y), int(m) if m else None, int(d) if d else None
    if not 1900 <= y <= date.today().year + 1:
        return ""
    try:
        date(y, m or 1, d or 1)
    except ValueError:
        return ""
    if precision == "year" or m is None:
        return f"{y:04d}"
    if precision == "month" or d is None:
        return f"{y:04d}-{m:02d}"
    return f"{y:04d}-{m:02d}-{d:02d}"


def fold(s: str) -> str:
    """Lowercase, accents removed: "Nestlé" -> "nestle"."""
    s = unicodedata.normalize("NFKD", s.casefold())
    return "".join(ch for ch in s if not unicodedata.combining(ch))


def safe_name(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", fold(s)).strip("-")
    return s[:80].rstrip("-") or "file"


def build_name(cfg: dict, cat: str, info: dict, ext: str) -> str:
    template = resolve(cfg, cat, info["subfolder"], "filename") or DEFAULT_FILENAME
    raw = template.format_map({k: info.get(k, "") for k in TEMPLATE_FIELDS} | {"title": info["title"] or "document"})
    return safe_name(raw) + ext


# ---- grouping (e.g. Employment/<employer>/<Payslips>) ---------------------

LEGAL_SUFFIX = re.compile(
    r"[\s,]+(s\.?a\.?s?\.?|s\.?à\.?r\.?l\.?|s\.?a\.?r\.?l\.?|sa|ag|gmbh|ltd\.?|inc\.?|llc|plc|bv|nv|se)$",
    re.IGNORECASE,
)


def clean_group(name: str) -> str:
    name = re.sub(r'[/\\:*?"<>|]', " ", name)
    name = re.sub(r"\s+", " ", name).strip(" .-")
    while True:  # "Acme Holding SA" -> "Acme Holding"
        stripped = LEGAL_SUFFIX.sub("", name).strip(" .,-")
        if stripped == name or not stripped:
            break
        name = stripped
    if name.islower():  # "acme" -> "Acme"
        name = name.title()
    return name[:60]


def match_group(name: str, known: list[str]) -> str:
    """Reuse an existing group when the new name is the same one written differently."""
    by_key = {fold(k): k for k in known}
    key = fold(name)
    if key in by_key:
        return by_key[key]
    for k, original in by_key.items():  # "Acme Switzerland" -> "Acme", and the reverse
        if key.startswith(k + " ") or k.startswith(key + " "):
            return original
    close = difflib.get_close_matches(key, list(by_key), n=1, cutoff=0.85)  # typos
    return by_key[close[0]] if close else name


def existing_groups(cfg: dict, cat: str) -> list[str]:
    """Group folders already on disk, so new files join them."""
    if not cfg["categories"][cat]["group_by"]:
        return []
    d = category_dir(cfg, cat)
    if not d.is_dir():
        return []
    return sorted(p.name for p in d.iterdir() if p.is_dir() and not p.name.startswith("."))


# ---- run files ------------------------------------------------------------

def new_run_file(kind: str) -> Path:
    """runs/<kind>-YYYYMMDD-HHMMSS.csv, with a -1, -2… suffix if that name is taken."""
    RUNS_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path, n = RUNS_DIR / f"{kind}-{stamp}.csv", 1
    while path.exists():
        path = RUNS_DIR / f"{kind}-{stamp}-{n}.csv"; n += 1
    return path


def latest_plan() -> Path | None:
    def key(p: Path):  # plan-DATE-TIME[-N]
        parts = p.stem.split("-")
        return parts[1], parts[2], int(parts[3]) if len(parts) > 3 else 0
    plans = sorted(RUNS_DIR.glob("plan-*.csv"), key=key) if RUNS_DIR.exists() else []
    return plans[-1] if plans else None


# ---- commands -------------------------------------------------------------

def plan(sources: list[Path], cfg: dict, refresh: bool):
    if cfg["root"] is None:
        sys.exit("Set 'destination' in config.toml or pass --into")
    Category, CategoryInPlace = make_category_model(cfg), make_category_model(cfg, folder_check=True)
    sorted_dirs = {category_dir(cfg, c) for c in cfg["categories"]}
    # (file, None) = new file; (file, category) = already-sorted file getting another pass
    items = [(f, None) for f in iter_files(sources, cfg, sorted_dirs)]
    items += list(iter_sorted(sources, cfg, sorted_dirs, refresh))
    finder = DuplicateFinder()
    index_sorted(sorted_dirs, finder)
    counts = {"move": 0, "keep": 0, "resort": 0, "unchanged": 0, "duplicate": 0, "skip": 0,
              "inbox": 0, "undated": 0}
    groups: dict[str, list[str]] = {}  # category -> group names seen on disk or in this run
    plan_path = new_run_file("plan")

    with open(plan_path, "x", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=PLAN_FIELDS)
        w.writeheader()
        for i, (f, sorted_cat) in enumerate(items, 1):
            print(f"[{i}/{len(items)}] {f}")
            row = {"src": f}
            if sorted_cat is None:  # sorted files are indexed already, so they'd match themselves
                try:
                    original = finder.check_and_add(f)
                except OSError as e:
                    w.writerow(row | {"action": "skip", "note": f"unreadable: {e}"}); counts["skip"] += 1
                    continue
                if original:
                    w.writerow(row | {"action": "duplicate", "note": f"same content as {original}"})
                    counts["duplicate"] += 1
                    continue

            folder = folder_label(f, sources)
            # a new file in a hand-made folder of the destination may stay there, if the folder fits it
            may_stay = sorted_cat is None and f.parent != cfg["root"] and f.is_relative_to(cfg["root"])
            fits = False
            try:
                text = extract_text(f, cfg)[: cfg["max_chars"]]
                if sorted_cat and sorted_cat != cfg["inbox"]:  # keep the category you already accepted
                    cat, suggested, confidence = sorted_cat, sorted_cat, ""
                else:  # new file, or an Inbox file getting a second chance
                    d = classify(f, text, folder, cfg, CategoryInPlace if may_stay else Category)
                    cat = d.category if d.confidence >= cfg["min_confidence"] else cfg["inbox"]
                    suggested, confidence = d.category, round(d.confidence, 2)
                    fits = getattr(d, "folder_fits", False)
                if not fits:
                    known = groups.setdefault(cat, existing_groups(cfg, cat))
                    info = describe(f, text, folder, cat, cfg, known)
                    if info["group"] and info["group"] not in known:
                        known.append(info["group"])
            except Exception as e:
                w.writerow(row | {"action": "skip", "note": f"model error: {e}"}); counts["skip"] += 1
                continue

            c = cfg["categories"][cat]
            ext = f.suffix.lower()
            if fits:  # dest = where it would go, in case you set the action to move
                stuck = ext in GOOGLE_STUBS and c["destination"]
                w.writerow(row | {"action": "keep", "category": cat, "suggested": suggested,
                                  "confidence": confidence, "dest": "" if stuck else category_dir(cfg, cat) / f.name,
                                  "note": "current folder fits: left in place"})
                counts["keep"] += 1
                continue
            name = f.name if ext in GOOGLE_STUBS else build_name(cfg, cat, info, ext)
            counts["inbox"] += cat == cfg["inbox"]
            counts["undated"] += not info["date"]
            row |= {"category": cat, "group": info["group"], "subfolder": info["subfolder"],
                    "date": info["date"], "suggested": suggested, "confidence": confidence}
            if sorted_cat:
                row["note"] = "already sorted: new place/name"

            if ext in GOOGLE_STUBS and c["destination"]:
                w.writerow(row | {"action": "skip", "note": "Google-native file can't leave Drive"})
                counts["skip"] += 1
                continue

            folder = category_dir(cfg, cat)
            for part in (info["group"], info["subfolder"]):
                if part:
                    folder = folder / part
            target = folder / name
            # name-2.pdf is where apply puts a file whose target name is taken, so it's right too
            if target == f or (f.parent == folder and f.suffix == target.suffix
                               and re.fullmatch(re.escape(target.stem) + r"-\d+", f.stem)):
                counts["unchanged"] += 1
                continue
            w.writerow(row | {"action": "move", "dest": target})
            counts["resort" if sorted_cat else "move"] += 1

    print(
        f"\nWrote {plan_path}:\n"
        f"  {counts['move']} new files to move ({counts['inbox']} to {cfg['inbox']}), "
        f"{counts['keep']} left in place (their folder fits)\n"
        f"  {counts['resort']} sorted files to move or rename, {counts['unchanged']} already right\n"
        f"  {counts['duplicate']} duplicates, {counts['skip']} skipped, "
        f"{counts['undated']} without a document date\n"
        "Review it, then run: uv run organize.py apply"
    )


def apply(plan_path: Path | None):
    plan_path = plan_path or latest_plan()
    if plan_path is None:
        sys.exit("No plan found in runs/. Run: uv run organize.py plan <folders>")
    if not plan_path.is_file():
        sys.exit(f"Plan not found: {plan_path}")
    print(f"Applying {plan_path}\n")

    moved, undo_path, undo_fh, log = 0, None, None, None
    try:
        with open(plan_path, newline="") as fh:
            for row in csv.DictReader(fh):
                if row["action"] != "move" or not row["dest"]:
                    continue
                src, dest = Path(row["src"]), Path(row["dest"])
                if not src.exists():
                    continue
                base, n = dest, 1
                while dest.exists() and dest != src:  # never overwrite
                    dest = base.with_stem(f"{base.stem}-{n}"); n += 1
                if dest == src:  # src already has the name or one of its -N variants
                    continue
                if log is None:  # create the undo log only once something moves
                    undo_path = new_run_file("undo")
                    undo_fh = open(undo_path, "x", newline="")
                    log = csv.writer(undo_fh)
                    log.writerow(["new_path", "original_path"])
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(src, dest)
                log.writerow([dest, src])
                undo_fh.flush()  # keep the log accurate even if the run is interrupted
                moved += 1
                print(f"{src}  ->  {dest}")
    finally:
        if undo_fh:
            undo_fh.close()

    if moved:
        print(f"\nMoved {moved} files. Undo log: {undo_path} (new path, original path).")
    else:
        print("Nothing to move: every 'move' row is already done or its source is gone.")


def main():
    ap = argparse.ArgumentParser(description="Sort files into folders with a local LLM.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("plan", help="classify files and write runs/plan-<timestamp>.csv (moves nothing)")
    p.add_argument("sources", nargs="+", help="folders to scan")
    p.add_argument("--into", help="destination root (overrides config.toml)")
    p.add_argument("--refresh", action="store_true",
                   help="also re-date and rename files already sorted (their category is kept; "
                        "Inbox files are re-classified)")
    a = sub.add_parser("apply", help="execute a plan (the latest one by default)")
    a.add_argument("plan", nargs="?", type=Path, help="plan file to apply (default: latest in runs/)")
    args = ap.parse_args()

    if args.cmd == "apply":
        return apply(args.plan)
    cfg = load_config()
    if args.into:
        cfg["root"] = expand(args.into)
    sources = [expand(s) for s in args.sources]
    for s in sources:
        if not s.is_dir():
            sys.exit(f"Not a folder: {s}")
    plan(sources, cfg, args.refresh)


if __name__ == "__main__":
    main()
